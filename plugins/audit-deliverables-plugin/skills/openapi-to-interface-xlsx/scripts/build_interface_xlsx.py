#!/usr/bin/env python3
"""Build per-중분류 Interface 설계서 xlsx files from OpenAPI JSON + summary xlsx.

See references/mapping.md for the full mapping spec.
"""
from __future__ import annotations

import argparse
import copy
import json
import re
import sys
from collections import defaultdict
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

import openpyxl
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet


# ────────────────────────────────────────────────────────────────────────────
# Config
# ────────────────────────────────────────────────────────────────────────────

TAG_ROLE_TO_SENDER = {
    "Buyer": "구매자",
    "Seller": "판매자",
    "Admin": "관리자",
    "Public": "외부 호출자",
    "Internal": "내부 서비스",
}

# Layout (column merges) per row-kind.
# Each entry is (start_col_letter, end_col_letter).
HEADER_LABEL_MERGES = [("A", "C"), ("D", "H"), ("I", "K"), ("L", "Q")]  # rows 2-4
URL_ROW_MERGES = [("A", "C"), ("D", "Q")]  # row 5
SECTION_LABEL_MERGES = [("A", "Q")]  # "Request Headers", "Request Parameters", etc.

REQ_HEADER_COLHEADER_MERGES = [
    ("B", "D"), ("E", "H"), ("I", "J"), ("K", "L"), ("N", "Q")
]
REQ_HEADER_DATAROW_MERGES = REQ_HEADER_COLHEADER_MERGES

REQ_PARAM_COLHEADER_MERGES = [
    ("B", "D"), ("E", "H"), ("I", "J"), ("K", "L"), ("N", "O"), ("P", "Q")
]
REQ_PARAM_DATAROW_MERGES = REQ_PARAM_COLHEADER_MERGES

RESP_PARAM_COLHEADER_MERGES = [
    ("B", "D"), ("E", "H"), ("I", "J"), ("K", "L"), ("N", "Q")
]
# Response data row merges depend on depth. depth=0 → B:D, depth=1 → C:D, depth>=2 → no name-merge.
def resp_data_merges_for_depth(depth: int) -> list[tuple[str, str]]:
    base = [("E", "H"), ("I", "J"), ("K", "L"), ("N", "Q")]
    if depth == 0:
        return [("B", "D"), *base]
    if depth == 1:
        return [("C", "D"), *base]
    return list(base)

DESCRIPTION_BODY_ROWS = 30  # rows for the bottom 설명 block

OPENAPI_TYPE_DISPLAY = {
    "string": "String",
    "integer": "Number",
    "number": "Number",
    "boolean": "Boolean",
    "array": "Array",
    "object": "Object",
}


# ────────────────────────────────────────────────────────────────────────────
# Data classes
# ────────────────────────────────────────────────────────────────────────────

@dataclass
class SummaryRow:
    no: int
    major: str            # 대분류
    middle: str           # 중분류 (raw, with newline)
    middle_clean: str     # for filename
    middle_prefix: str    # IF-XXX-YYY
    if_id: str            # IF-DTS-MDB-10 (summary xlsx format)
    if_name: str          # interface 명
    description: str
    remark: str


@dataclass
class Operation:
    if_id_normalized: tuple[str, int]   # ("IF-DTS-MDB", 10)
    if_id_json: str                     # IF-DTS-MDB-010 (as in summary)
    if_name_from_summary: str           # name part after "[IF-…]"
    method: str
    path: str
    tags: list[str]
    op: dict
    info_title: str
    servers: list[dict]
    schemas: dict


# ────────────────────────────────────────────────────────────────────────────
# Style helpers
# ────────────────────────────────────────────────────────────────────────────

class StylePack:
    """Holds copyable style references extracted from the template sheet."""

    def __init__(self, template_ws: Worksheet):
        self.template = template_ws
        # Cache: (row, col_letter) → cell
        self.cache: dict[tuple[int, str], Any] = {}

    def apply(self, target_cell, src_row: int, src_col_letter: str):
        src = self.template[f"{src_col_letter}{src_row}"]
        if src.has_style:
            target_cell.font = copy.copy(src.font)
            target_cell.fill = copy.copy(src.fill)
            target_cell.border = copy.copy(src.border)
            target_cell.alignment = copy.copy(src.alignment)
            target_cell.number_format = src.number_format
            target_cell.protection = copy.copy(src.protection)


# ────────────────────────────────────────────────────────────────────────────
# Loading
# ────────────────────────────────────────────────────────────────────────────

ID_RE = re.compile(r"\[(IF-[A-Z0-9-]+)\]")
NUM_TAIL_RE = re.compile(r"(.+-)(\d+)$")


def normalize_if_id(raw_id: str) -> tuple[str, int]:
    m = NUM_TAIL_RE.match(raw_id)
    if not m:
        return (raw_id, -1)
    prefix = m.group(1).rstrip("-")
    num = int(m.group(2))
    return (prefix, num)


def sanitize_filename(name: str) -> str:
    name = name.replace("\n", " ").replace("\r", " ")
    for ch in '\\/:*?"<>|':
        name = name.replace(ch, "")
    return name.strip()


def sanitize_sheetname(name: str) -> str:
    name = sanitize_filename(name)
    return name[:31]


def load_summary(path: Path) -> list[SummaryRow]:
    wb = openpyxl.load_workbook(path, data_only=True)
    sheet_name = "요약표" if "요약표" in wb.sheetnames else wb.sheetnames[0]
    ws = wb[sheet_name]
    rows: list[SummaryRow] = []
    current_major = ""
    current_middle = ""
    for row_idx, row in enumerate(ws.iter_rows(values_only=True), 1):
        if row_idx < 3:  # rows 1-2 are blank + header
            continue
        no, major, middle, if_id, if_name, desc, remark = (row + (None,) * 7)[:7]
        if major:
            current_major = str(major)
        if middle:
            current_middle = str(middle)
        if not if_id:
            continue
        middle_clean = re.sub(r"\s+", " ", current_middle.replace("\n", " ")).strip()
        # extract prefix from if_id (everything except trailing -N)
        norm = normalize_if_id(str(if_id).strip())
        prefix = norm[0]
        rows.append(SummaryRow(
            no=int(no) if no is not None else 0,
            major=current_major,
            middle=current_middle,
            middle_clean=middle_clean,
            middle_prefix=prefix,
            if_id=str(if_id).strip(),
            if_name=str(if_name or "").strip(),
            description=str(desc or "").strip(),
            remark=str(remark or "").strip(),
        ))
    return rows


def load_openapi_files(paths: list[Path]) -> list[Operation]:
    ops: list[Operation] = []
    for p in paths:
        with p.open(encoding="utf-8") as f:
            doc = json.load(f)
        info_title = doc.get("info", {}).get("title", "")
        servers = doc.get("servers", [])
        schemas = doc.get("components", {}).get("schemas", {})
        for path, methods in doc.get("paths", {}).items():
            for method, op in methods.items():
                if method.lower() not in {"get", "post", "put", "delete", "patch", "head", "options"}:
                    continue
                summary = op.get("summary", "") or ""
                m = ID_RE.search(summary)
                if not m:
                    continue
                raw_id = m.group(1)
                norm = normalize_if_id(raw_id)
                name_part = ID_RE.sub("", summary).strip()
                ops.append(Operation(
                    if_id_normalized=norm,
                    if_id_json=raw_id,
                    if_name_from_summary=name_part,
                    method=method.upper(),
                    path=path,
                    tags=op.get("tags", []),
                    op=op,
                    info_title=info_title,
                    servers=servers,
                    schemas=schemas,
                ))
    return ops


# ────────────────────────────────────────────────────────────────────────────
# OpenAPI helpers
# ────────────────────────────────────────────────────────────────────────────

def resolve_ref(schema: dict, schemas: dict, _stack: set | None = None) -> dict:
    """Resolve $ref. Returns the resolved schema (not deep-copied)."""
    if _stack is None:
        _stack = set()
    if not isinstance(schema, dict):
        return {}
    ref = schema.get("$ref")
    if not ref:
        return schema
    name = ref.rsplit("/", 1)[-1]
    if name in _stack:
        return {}
    if name not in schemas:
        return {}
    _stack.add(name)
    resolved = resolve_ref(schemas[name], schemas, _stack)
    # Merge: ref-pointer's siblings (description, etc.) override
    merged = dict(resolved)
    for k, v in schema.items():
        if k != "$ref":
            merged[k] = v
    return merged


def schema_type_display(schema: dict) -> str:
    t = (schema.get("type") or "").lower()
    if t in OPENAPI_TYPE_DISPLAY:
        return OPENAPI_TYPE_DISPLAY[t]
    # if no type but has properties, treat as object
    if "properties" in schema:
        return "Object"
    if "items" in schema:
        return "Array"
    if "enum" in schema:
        return "String"
    return ""


def schema_data_field(schema: dict) -> str:
    """Format-only info for the 데이터(M) column; no specific examples."""
    parts: list[str] = []
    if "enum" in schema and schema["enum"]:
        parts.append(", ".join(str(v) for v in schema["enum"]))
    fmt = schema.get("format")
    if fmt:
        parts.append(fmt)
    return " / ".join(parts) if parts else ""


def schema_size(schema: dict) -> str:
    if "maxLength" in schema:
        return str(schema["maxLength"])
    return "가변"


def extract_sender_from_tags(tags: list[str]) -> str:
    for tag in tags:
        m = re.search(r"\[([A-Za-z]+)\]", tag)
        if m and m.group(1) in TAG_ROLE_TO_SENDER:
            return TAG_ROLE_TO_SENDER[m.group(1)]
    return ""


def extract_receiver_from_title(title: str) -> str:
    """Map OpenAPI info.title to a Korean system name."""
    mapping = {
        "Marketplace": "마켓플레이스",
        "Data Catalog": "데이터 카탈로그",
        "Data Quality": "품질검사",
        "DTS": "데이터유통시스템",
    }
    for keyword, ko in mapping.items():
        if keyword.lower() in title.lower():
            return ko
    return title  # fallback to title verbatim


def build_url(method: str, path: str, servers: list[dict]) -> str:
    server_url = ""
    if servers:
        server_url = servers[0].get("url", "")
        if "localhost" in server_url or "127.0.0.1" in server_url:
            server_url = "{API}"
    return f"{method} {server_url}{path}".strip()


# ────────────────────────────────────────────────────────────────────────────
# Parameter extraction
# ────────────────────────────────────────────────────────────────────────────

@dataclass
class FlatParam:
    name: str
    description: str
    type_display: str
    size: str
    data: str
    required: bool
    location: str   # path / query / header / body
    default: Any = None


def extract_request_headers(op: dict, schemas: dict) -> list[FlatParam]:
    """Request Headers 섹션: HTTP header + path variable.

    path variable 은 URL 의 일부지만 호출 측 입력값이라는 점에서 헤더 영역에 같이 모아 표기한다.
    """
    out: list[FlatParam] = []
    for p in op.get("parameters", []) or []:
        loc = (p.get("in") or "").lower()
        if loc not in {"header", "path"}:
            continue
        schema = resolve_ref(p.get("schema") or {}, schemas)
        out.append(FlatParam(
            name=p.get("name", ""),
            description=p.get("description", "") or "",
            type_display=schema_type_display(schema),
            size=schema_size(schema),
            data=schema_data_field(schema),
            required=bool(p.get("required")),
            location=loc,
            default=schema.get("default"),
        ))
    return out


def extract_request_parameters(op: dict, schemas: dict) -> list[FlatParam]:
    out: list[FlatParam] = []
    for p in op.get("parameters", []) or []:
        loc = (p.get("in") or "").lower()
        if loc != "query":
            continue
        schema = resolve_ref(p.get("schema") or {}, schemas)
        out.append(FlatParam(
            name=p.get("name", ""),
            description=p.get("description", "") or "",
            type_display=schema_type_display(schema),
            size=schema_size(schema),
            data=schema_data_field(schema),
            required=bool(p.get("required")),
            location=loc,
            default=schema.get("default"),
        ))
    # requestBody: flatten first media-type schema's properties (depth 0 only)
    body = op.get("requestBody") or {}
    content = body.get("content") or {}
    if content:
        media_key = next(iter(content))
        media = content[media_key]
        body_schema = resolve_ref(media.get("schema") or {}, schemas)
        required_set = set(body_schema.get("required", []))
        for name, ps in (body_schema.get("properties") or {}).items():
            ps_resolved = resolve_ref(ps, schemas)
            out.append(FlatParam(
                name=name,
                description=ps_resolved.get("description", "") or "",
                type_display=schema_type_display(ps_resolved),
                size=schema_size(ps_resolved),
                data=schema_data_field(ps_resolved),
                required=name in required_set,
                location="body",
                default=ps_resolved.get("default"),
            ))
    return out


# ────────────────────────────────────────────────────────────────────────────
# Response tree
# ────────────────────────────────────────────────────────────────────────────

@dataclass
class RespRow:
    a_label: str           # "1" / "D1" / "D1-3"
    depth: int             # 0 = top-level, 1+ deeper
    name: str
    description: str
    type_display: str
    size: str
    data: str


def flatten_response(schema: dict, schemas: dict) -> list[RespRow]:
    """Flatten a response schema using the D{n} / D{n}-{seq} tree convention.

    Top-level properties get number 1, 2, 3 …; their descendants are emitted
    in document order. Before any descent into deeper levels, direct children
    are labeled "D{n}". As soon as we descend (or once we've emitted any
    non-direct-child item under #n), all subsequent rows under #n are labeled
    "D{n}-{seq}" with seq starting at 1 and increasing per row.
    """
    out: list[RespRow] = []
    schema = resolve_ref(schema, schemas)
    top_props: dict[str, dict] = schema.get("properties") or {}

    for idx, (name, ps) in enumerate(top_props.items(), 1):
        ps_resolved = resolve_ref(ps, schemas)
        out.append(RespRow(
            a_label=str(idx),
            depth=0,
            name=name,
            description=ps_resolved.get("description", "") or "",
            type_display=schema_type_display(ps_resolved),
            size=schema_size(ps_resolved),
            data=schema_data_field(ps_resolved),
        ))
        # Recurse into this top-level item
        state = {"seq": 0, "descended": False}
        _walk_top_subtree(ps_resolved, schemas, top_idx=idx, depth=1, out=out, state=state)
    return out


def _walk_top_subtree(node: dict, schemas: dict, *, top_idx: int, depth: int,
                       out: list[RespRow], state: dict):
    """Walk the children of a top-level item, emitting rows.

    state['seq'] tracks the running descendant counter for the top item.
    state['descended'] becomes True once we've gone past depth 1.
    """
    # Unwrap array items
    if (node.get("type") == "array") or ("items" in node):
        items_schema = resolve_ref(node.get("items") or {}, schemas)
        # The array's element fields are at the same depth as direct children
        node = items_schema

    props: dict[str, dict] = node.get("properties") or {}
    if not props:
        return

    # Track whether any deeper recursion has happened in this top subtree
    for name, ps in props.items():
        ps_resolved = resolve_ref(ps, schemas)
        if state["descended"]:
            state["seq"] += 1
            label = f"D{top_idx}-{state['seq']}"
        else:
            label = f"D{top_idx}"
        out.append(RespRow(
            a_label=label,
            depth=depth,
            name=name,
            description=ps_resolved.get("description", "") or "",
            type_display=schema_type_display(ps_resolved),
            size=schema_size(ps_resolved),
            data=schema_data_field(ps_resolved),
        ))
        # Recurse if this property has a sub-structure
        has_children = False
        if ps_resolved.get("type") == "object" or ps_resolved.get("properties"):
            has_children = True
        if ps_resolved.get("type") == "array" or "items" in ps_resolved:
            items_schema = resolve_ref(ps_resolved.get("items") or {}, schemas)
            if items_schema.get("properties"):
                has_children = True
        if has_children:
            state["descended"] = True
            _walk_top_subtree(ps_resolved, schemas, top_idx=top_idx,
                              depth=depth + 1, out=out, state=state)


# ────────────────────────────────────────────────────────────────────────────
# Example response synthesizer
# ────────────────────────────────────────────────────────────────────────────

def synthesize_example(schema: dict, schemas: dict, _depth: int = 0) -> Any:
    if _depth > 6:
        return None
    schema = resolve_ref(schema, schemas)
    if "example" in schema:
        return schema["example"]
    if "enum" in schema and schema["enum"]:
        return schema["enum"][0]
    t = (schema.get("type") or "").lower()
    if t == "object" or "properties" in schema:
        return {
            k: synthesize_example(v, schemas, _depth + 1)
            for k, v in (schema.get("properties") or {}).items()
        }
    if t == "array" or "items" in schema:
        return [synthesize_example(schema.get("items") or {}, schemas, _depth + 1)]
    fmt = schema.get("format", "")
    if t == "string":
        if fmt == "date":
            return "YYYY-MM-DD"
        if fmt == "date-time":
            return "YYYY-MM-DDTHH:MM:SSZ"
        return ""
    if t == "integer" or t == "number":
        return 0
    if t == "boolean":
        return False
    return None


# ────────────────────────────────────────────────────────────────────────────
# Sheet building
# ────────────────────────────────────────────────────────────────────────────

def write_row(ws: Worksheet, row_idx: int, values: dict[str, Any],
              merges: list[tuple[str, str]], style: StylePack,
              style_src_row: int, height: float | None = None):
    """Write cells and apply merges + styles. values keyed by column letter."""
    if height:
        ws.row_dimensions[row_idx].height = height
    # Determine the set of columns we touch
    all_cols = set()
    for s, e in merges:
        s_i, e_i = openpyxl.utils.column_index_from_string(s), openpyxl.utils.column_index_from_string(e)
        for c in range(s_i, e_i + 1):
            all_cols.add(get_column_letter(c))
    # Always cover A..Q for styling
    for c in range(1, 18):  # A..Q
        all_cols.add(get_column_letter(c))

    for col_letter in all_cols:
        cell = ws[f"{col_letter}{row_idx}"]
        style.apply(cell, style_src_row, col_letter)
        if col_letter in values:
            cell.value = values[col_letter]
    for s, e in merges:
        if s == e:
            continue
        ws.merge_cells(f"{s}{row_idx}:{e}{row_idx}")


def copy_column_widths(template_ws: Worksheet, target_ws: Worksheet):
    for col_letter, dim in template_ws.column_dimensions.items():
        if dim.width:
            target_ws.column_dimensions[col_letter].width = dim.width


def build_interface_sheet(wb: openpyxl.Workbook, sheet_name: str,
                          summary_row: SummaryRow | None,
                          operation: Operation | None,
                          template_ws: Worksheet,
                          style: StylePack):
    ws = wb.create_sheet(sheet_name)
    copy_column_widths(template_ws, ws)

    # ── Header (rows 2-5) ──────────────────────────────────────────────────
    # row 1 spacer (full-width merge)
    write_row(ws, 1, {}, [("A", "Q")], style, 1, height=template_ws.row_dimensions[1].height)

    if_id = summary_row.if_id if summary_row else (operation.if_id_json if operation else "")
    if_name = (summary_row.if_name if summary_row and summary_row.if_name
               else (operation.if_name_from_summary if operation else ""))
    sender = extract_sender_from_tags(operation.tags) if operation else ""
    receiver = extract_receiver_from_title(operation.info_title) if operation else ""
    url = build_url(operation.method, operation.path, operation.servers) if operation else ""

    write_row(ws, 2, {"A": "인터페이스ID", "D": if_id, "I": "인터페이스명", "L": if_name},
              [("A", "C"), ("D", "H"), ("I", "K"), ("L", "Q")], style, 2,
              height=template_ws.row_dimensions[2].height)
    write_row(ws, 3, {"A": "데이터 송신처", "D": sender, "I": "데이터 수신처", "L": receiver},
              [("A", "C"), ("D", "H"), ("I", "K"), ("L", "Q")], style, 3,
              height=template_ws.row_dimensions[3].height)
    write_row(ws, 4, {"A": "프로토콜", "D": "HTTP", "I": "발생주기", "L": "수동호출"},
              [("A", "C"), ("D", "H"), ("I", "K"), ("L", "Q")], style, 4,
              height=template_ws.row_dimensions[4].height)
    write_row(ws, 5, {"A": "URL", "D": url}, URL_ROW_MERGES, style, 5,
              height=template_ws.row_dimensions[5].height)

    cursor = 6

    # ── Request Headers ────────────────────────────────────────────────────
    write_row(ws, cursor, {"A": "Request Headers"}, SECTION_LABEL_MERGES, style, 6,
              height=template_ws.row_dimensions[6].height)
    cursor += 1
    write_row(ws, cursor, {
        "A": "번호", "B": "데이터명", "E": "설명", "I": "형태",
        "K": "크기", "M": "데이터", "N": "비고"
    }, REQ_HEADER_COLHEADER_MERGES, style, 7,
              height=template_ws.row_dimensions[7].height)
    cursor += 1

    req_headers = extract_request_headers(operation.op, operation.schemas) if operation else []
    if not req_headers:
        write_row(ws, cursor, {}, REQ_HEADER_DATAROW_MERGES, style, 8,
                  height=template_ws.row_dimensions[8].height or 20.0)
        cursor += 1
    else:
        for idx, p in enumerate(req_headers, 1):
            remark_bits = [p.location] if p.location else []
            if p.required:
                remark_bits.append("필수")
            write_row(ws, cursor, {
                "A": idx, "B": p.name, "E": p.description,
                "I": p.type_display, "K": p.size, "M": p.data,
                "N": ", ".join(remark_bits)
            }, REQ_HEADER_DATAROW_MERGES, style, 8,
                      height=template_ws.row_dimensions[8].height or 20.0)
            cursor += 1

    # ── Request Parameters ────────────────────────────────────────────────
    write_row(ws, cursor, {"A": "Request Parameters"}, SECTION_LABEL_MERGES, style, 9,
              height=template_ws.row_dimensions[9].height)
    cursor += 1
    write_row(ws, cursor, {
        "A": "번호", "B": "데이터명", "E": "설명", "I": "형태",
        "K": "크기", "M": "데이터", "N": "필수", "P": "비고"
    }, REQ_PARAM_COLHEADER_MERGES, style, 10,
              height=template_ws.row_dimensions[10].height)
    cursor += 1

    req_params = extract_request_parameters(operation.op, operation.schemas) if operation else []
    if not req_params:
        write_row(ws, cursor, {}, REQ_PARAM_DATAROW_MERGES, style, 11,
                  height=template_ws.row_dimensions[11].height or 20.0)
        cursor += 1
    else:
        for idx, p in enumerate(req_params, 1):
            remark_bits = [p.location] if p.location else []
            if p.default is not None:
                remark_bits.append(f"기본값: {p.default}")
            write_row(ws, cursor, {
                "A": idx, "B": p.name, "E": p.description,
                "I": p.type_display, "K": p.size, "M": p.data,
                "N": ("Y" if p.required else "N"),
                "P": ", ".join(remark_bits)
            }, REQ_PARAM_DATAROW_MERGES, style, 11,
                      height=template_ws.row_dimensions[11].height or 20.0)
            cursor += 1

    # ── Response Parameters ───────────────────────────────────────────────
    write_row(ws, cursor, {"A": "Response Parameters"}, SECTION_LABEL_MERGES, style, 13,
              height=template_ws.row_dimensions[13].height or 20.0)
    cursor += 1
    write_row(ws, cursor, {
        "A": "번호", "B": "데이터명", "E": "설명", "I": "형태",
        "K": "크기", "M": "데이터", "N": "비고"
    }, RESP_PARAM_COLHEADER_MERGES, style, 14,
              height=template_ws.row_dimensions[14].height or 20.0)
    cursor += 1

    resp_rows: list[RespRow] = []
    if operation:
        responses = operation.op.get("responses") or {}
        # prefer 200, else first
        candidate = responses.get("200") or next(iter(responses.values()), {})
        content = (candidate or {}).get("content") or {}
        media = content.get("application/json") or (next(iter(content.values())) if content else {})
        schema = (media or {}).get("schema") or {}
        if schema:
            resp_rows = flatten_response(schema, operation.schemas)

    if not resp_rows:
        write_row(ws, cursor, {}, resp_data_merges_for_depth(0), style, 15,
                  height=template_ws.row_dimensions[15].height or 20.0)
        cursor += 1
    else:
        for r in resp_rows:
            name_col = {0: "B", 1: "C"}.get(r.depth, "D")
            merges = resp_data_merges_for_depth(r.depth)
            write_row(ws, cursor, {
                "A": r.a_label, name_col: r.name, "E": r.description,
                "I": r.type_display, "K": r.size, "M": r.data
            }, merges, style, 15,
                      height=template_ws.row_dimensions[15].height or 20.0)
            cursor += 1

    # ── 설명 ──────────────────────────────────────────────────────────────
    write_row(ws, cursor, {"A": "설명"}, SECTION_LABEL_MERGES, style, 18,
              height=template_ws.row_dimensions[18].height or 20.0)
    cursor += 1

    description_text = ""
    if operation:
        responses = operation.op.get("responses") or {}
        candidate = responses.get("200") or next(iter(responses.values()), {})
        content = (candidate or {}).get("content") or {}
        media = content.get("application/json") or (next(iter(content.values())) if content else {})
        schema = (media or {}).get("schema") or {}
        example = synthesize_example(schema, operation.schemas) if schema else None
        example_json = json.dumps(example, ensure_ascii=False, indent=2) if example is not None else "{}"
        description_text = (
            "서비스 리턴값: Response OK(200), BadRequest(400)\n\n"
            f"예시 응답:\n{example_json}"
        )

    body_height = DESCRIPTION_BODY_ROWS
    # write first row of description
    write_row(ws, cursor, {"A": description_text}, [("A", "Q")], style, 19,
              height=None)
    ws.row_dimensions[cursor].height = 200
    # extend merge over DESCRIPTION_BODY_ROWS rows
    end_row = cursor + body_height - 1
    if end_row > cursor:
        ws.unmerge_cells(f"A{cursor}:Q{cursor}")
        ws.merge_cells(f"A{cursor}:Q{end_row}")
    cursor = end_row + 1

    return ws


def build_summary_sheet(wb: openpyxl.Workbook, group_rows: list[SummaryRow],
                        op_index: dict[tuple[str, int], Operation]):
    ws = wb.create_sheet("요약", 0)
    ws.append(["Interface ID", "Interface 명", "Method/URL", "설명"])
    for r in group_rows:
        op = op_index.get(normalize_if_id(r.if_id))
        method_url = f"{op.method} {op.path}" if op else "(미구현)"
        ws.append([r.if_id, r.if_name, method_url, r.description])
    # auto width-ish
    widths = [22, 35, 60, 80]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


# ────────────────────────────────────────────────────────────────────────────
# Main
# ────────────────────────────────────────────────────────────────────────────

def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--summary", required=True, type=Path)
    ap.add_argument("--template", required=True, type=Path)
    ap.add_argument("--json", required=True, nargs="+", type=Path)
    ap.add_argument("--out", type=Path,
                    help="중분류별 분리 xlsx 출력 디렉토리. --append-to 모드면 생략 가능.")
    ap.add_argument("--append-to", type=Path,
                    help="이 워크북에 IF 시트들을 append. summary xlsx 의 ID 순서대로 정렬. 기존 시트는 보존.")
    args = ap.parse_args()
    if not args.out and not args.append_to:
        ap.error("--out 또는 --append-to 중 하나는 필요.")

    summary_rows = load_summary(args.summary)
    operations = load_openapi_files(args.json)

    op_index: dict[tuple[str, int], Operation] = {}
    for op in operations:
        op_index[op.if_id_normalized] = op

    template_wb = openpyxl.load_workbook(args.template)
    template_ws = template_wb["템플릿"] if "템플릿" in template_wb.sheetnames else template_wb.worksheets[0]
    style = StylePack(template_ws)

    used_op_keys: set[tuple[str, int]] = set()

    # ── Append-mode: 단일 워크북에 summary 순서대로 IF 시트만 추가 ──────────────
    if args.append_to:
        wb = openpyxl.load_workbook(args.append_to)
        existing = set(wb.sheetnames)
        for r in summary_rows:
            sheet_name = sanitize_sheetname(r.if_id)
            if sheet_name in existing:
                # 이미 같은 이름의 시트가 있으면 제거 후 재생성
                del wb[sheet_name]
            key = normalize_if_id(r.if_id)
            op = op_index.get(key)
            if not op:
                sheet_name = sanitize_sheetname(r.if_id + " (미구현)")
                if sheet_name in wb.sheetnames:
                    del wb[sheet_name]
            build_interface_sheet(wb, sheet_name, r, op, template_ws, style)
            if op:
                used_op_keys.add(key)
        wb.save(args.append_to)
        print(f"appended {len(summary_rows)} sheets to {args.append_to}")
        unmapped = [op for key, op in op_index.items() if key not in used_op_keys]
        if unmapped:
            print(f"warning: {len(unmapped)} operations in JSON were not matched to summary 요약표:")
            for op in unmapped:
                print(f"  - {op.if_id_json}  {op.method} {op.path}")
        return

    # ── 기본 모드: 중분류별 파일 분리 ────────────────────────────────────────
    # Group summary rows by middle prefix
    groups: dict[str, tuple[str, list[SummaryRow]]] = {}
    for r in summary_rows:
        if r.middle_prefix not in groups:
            groups[r.middle_prefix] = (r.middle_clean, [])
        groups[r.middle_prefix][1].append(r)

    args.out.mkdir(parents=True, exist_ok=True)

    for prefix, (middle_name, rows) in groups.items():
        wb = openpyxl.Workbook()
        # remove default sheet
        wb.remove(wb.active)
        # build per-IF sheets
        for r in rows:
            key = normalize_if_id(r.if_id)
            op = op_index.get(key)
            sheet_name = sanitize_sheetname(r.if_id + ("" if op else " (미구현)"))
            build_interface_sheet(wb, sheet_name, r, op, template_ws, style)
            if op:
                used_op_keys.add(key)
        # build summary sheet on top
        build_summary_sheet(wb, rows, op_index)

        fname = sanitize_filename(middle_name) + ".xlsx"
        wb.save(args.out / fname)
        print(f"wrote {args.out / fname}  ({len(rows)} interfaces)")

    # _unmapped: operations in JSON but not in summary
    unmapped = [op for key, op in op_index.items() if key not in used_op_keys]
    if unmapped:
        wb = openpyxl.Workbook()
        wb.remove(wb.active)
        for op in unmapped:
            sheet_name = sanitize_sheetname(op.if_id_json)
            build_interface_sheet(wb, sheet_name, None, op, template_ws, style)
        wb.save(args.out / "_unmapped.xlsx")
        print(f"wrote {args.out / '_unmapped.xlsx'}  ({len(unmapped)} unmapped)")


if __name__ == "__main__":
    main()
